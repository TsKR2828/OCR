"""Excel Renderer — 從 timeline segments 產出 transcript.xlsx."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

from ..schema import Segment, fmt_ts


HEADERS = [
    "時間戳",
    "章節",
    "角色",
    "遊戲原文(OCR)",
    "語音辨識(ASR)",
    "說話者",
    "聊天室密度",
    "SC",
    "實況主反應",
    "一致性",
    "精彩度",
]

COL_WIDTHS = [12, 18, 10, 35, 35, 10, 14, 6, 10, 14, 10]

SPIKE_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
CONFLICT_FILL = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")


def render_excel(segments: list[Segment], output_path: Path, title: str = "") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] if title else "transcript"

    header_font = Font(bold=True)
    for col_idx, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = width

    ws.auto_filter.ref = f"A1:K1"

    for row_idx, seg in enumerate(segments, 2):
        ts = fmt_ts(seg.time_start)
        chapter = seg.ocr.chapter if seg.ocr else ""
        character = seg.ocr.character if seg.ocr else ""
        ocr_text = seg.ocr.dialogue if seg.ocr else ""
        asr_text = seg.asr.text if seg.asr else ""
        speaker = seg.asr.speaker_guess if seg.asr else ""

        chat_density = ""
        sc_count = ""
        if seg.chat:
            chat_density = f"{seg.chat.density_per_min:.0f}/min"
            if seg.chat.superchat_count > 0:
                sc_count = str(seg.chat.superchat_count)

        reaction = "Y" if seg.events.streamer_reaction or seg.events.chat_spike else ""
        match_status = seg.merge.match_status
        score = seg.score.total

        row_data = [
            ts, chapter, character, ocr_text, asr_text,
            speaker, chat_density, sc_count, reaction, match_status,
            score if score > 0 else "",
        ]

        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

        if seg.events.chat_spike:
            for col_idx in range(1, len(HEADERS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = SPIKE_FILL

        if match_status == "conflict":
            for col_idx in range(1, len(HEADERS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = CONFLICT_FILL

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"[Excel] 寫出 → {output_path} ({len(segments)} 行)")
